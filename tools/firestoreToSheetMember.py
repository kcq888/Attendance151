from __future__ import print_function
import os
from datetime import datetime
import pytz    # $ pip install pytz
import tzlocal # $ pip install tzlocal
# google sheet
import pickle
from googleapiclient.discovery import build
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
# Firebase Imports
import firebase_admin
from firebase_admin import credentials
from firebase_admin import firestore
from openpyxl import Workbook
from openpyxl.styles import Font

class FireStoreToSheetMember():
    SCOPES = ['https://www.googleapis.com/auth/spreadsheets', "https://www.googleapis.com/auth/drive.file", "https://www.googleapis.com/auth/drive"]
    SHEET_ID = '1HU3o-d53u1tE9vZLjnlCFDJ1NwlDU_gf4dgwLvrY-Ck'
    Season = 'Season2025-2026'
    Members = 'members'
    rfids = 'rfids'

    def __init__(self):
        # Firestore setup
        self.fcred = credentials.Certificate(os.environ.get("GOOGLE_APPLICATION_CREDENTIALS"))
        firebase_admin.initialize_app(self.fcred)
        self.db = firestore.client()
        
        # Google Sheet API setup
        self.gscreds = None
        if os.path.exists('token.pickle'):
            with open('token.pickle', 'rb') as token:
                self.gscreds = pickle.load(token)
        self.service = build('sheets', 'v4', credentials=self.gscreds)
        if not self.gscreds or not self.gscreds.valid:
            if self.gscreds and self.gscreds.expired and self.gscreds.refresh_token:
                self.gscreds.refresh(Request())
            else:
                flow = InstalledAppFlow.from_client_secrets_file(
                    'credentials.json', self.SCOPES)
                self.gscreds = flow.run_local_server(port=0)
            with open('token.pickle', 'wb') as token:
                pickle.dump(self.gscreds, token)
        self.index = 2
        self.name = 'A'
        self.rfid = 'B'
        self.value_input_option = 'USER_ENTERED'
        self.insert_data_option = 'OVERWRITE'

    def getMembers(self):
        wb = self.setupWorkbook()
        ws = wb.active
        colref = self.db.collection(self.Season, self.Members, self.rfids)
        rfids = colref.list_documents()

        for rfid in rfids:
            print(rfid.id)
            if rfid is not None:
                data = rfid.get()
                ws.append([data.get('First') + ' ' + data.get('Last'), rfid.id])

        membersheet = os.getcwd() + "/Members.xlsx"
        wb.save(membersheet)
        
    def setupWorkbook(self):
        print("setupWorkbook")
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendant"
        ws.append(["Name", "RFID"])
        return wb

if __name__ == "__main__":
    firestoreToSheetMember = FireStoreToSheetMember()
    firestoreToSheetMember.getMembers()
    